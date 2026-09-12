"""
Edge-case coverage for app/importer/excel_importer.py - the part that
decides how ambiguous spreadsheet cells get turned into dates, and which
rows/cells get flagged for human review. These are pure functions (no DB),
so they run without Postgres - `pytest` from backend/ is enough.

Deliberately test against the exact messy patterns seen in the customer's
real workbook (copy-forward P==A, "Auto Phospating Plant" with no M/c No.,
week codes, day-of-month, mixed date formats) rather than only clean
synthetic input.
"""

from datetime import date, datetime

import openpyxl
import pytest

from app.importer.excel_importer import (
    classify_cell_value,
    _parse_exact_date,
    week_code_to_date,
    _resolve_financial_year_calendar_year,
    is_footer_row,
    normalize_for_dedup,
    parse_pm_sheet,
)


# ---------------------------------------------------------------------------
# classify_cell_value
# ---------------------------------------------------------------------------

class TestClassifyCellValue:
    def test_empty_variants(self):
        assert classify_cell_value(None) == ("EMPTY", None)
        assert classify_cell_value("") == ("EMPTY", None)

    def test_exact_date_string_dot_separated(self):
        assert classify_cell_value("8.04.2023") == ("EXACT_DATE", "8.04.2023")

    def test_exact_date_string_slash_separated(self):
        t, raw = classify_cell_value("08/04/2023")
        assert t == "EXACT_DATE"

    def test_exact_date_datetime_object(self):
        t, raw = classify_cell_value(datetime(2023, 4, 8))
        assert t == "EXACT_DATE"
        assert raw == "08.04.2023"

    def test_exact_date_date_object(self):
        t, raw = classify_cell_value(date(2023, 4, 8))
        assert t == "EXACT_DATE"

    @pytest.mark.parametrize("code", ["W1", "w1", "W5", "W3"])
    def test_week_code(self, code):
        assert classify_cell_value(code) == ("WEEK_CODE", code.upper())

    def test_week_code_w6_is_not_valid(self):
        # only W1-W5 are locked business rule; W6 has no meaning
        t, raw = classify_cell_value("W6")
        assert t == "UNKNOWN"

    def test_day_of_month(self):
        assert classify_cell_value("19") == ("DAY_OF_MONTH", "19")
        assert classify_cell_value(" 5 ") == ("DAY_OF_MONTH", "5")

    def test_unknown_free_text(self):
        t, raw = classify_cell_value("done - see remarks")
        assert t == "UNKNOWN"
        assert raw == "done - see remarks"

    def test_unknown_never_silently_guessed_as_a_date(self):
        # "NA", "-", "pending" etc must never resolve to a date
        for junk in ["NA", "-", "pending", "TBD", "??"]:
            t, _ = classify_cell_value(junk)
            assert t == "UNKNOWN", f"{junk!r} should be UNKNOWN, not guessed"


# ---------------------------------------------------------------------------
# date parsing helpers
# ---------------------------------------------------------------------------

class TestParseExactDate:
    def test_two_digit_year_assumes_2000s(self):
        assert _parse_exact_date("8.04.23") == date(2023, 4, 8)

    def test_four_digit_year(self):
        assert _parse_exact_date("8.04.2023") == date(2023, 4, 8)

    def test_invalid_calendar_date_returns_none(self):
        # Feb 30 doesn't exist - must not raise, must not silently clamp
        assert _parse_exact_date("30.02.2023") is None

    def test_garbage_returns_none(self):
        assert _parse_exact_date("not a date") is None


class TestWeekCodeToDate:
    def test_locked_convention_default(self):
        # W1=1-7 (day1), W2=8-14(day8), W3=15-21(day15), W4=22-28(day22), W5=29-end
        assert week_code_to_date("W1", "April", 2023) == date(2023, 4, 1)
        assert week_code_to_date("W2", "April", 2023) == date(2023, 4, 8)
        assert week_code_to_date("W3", "April", 2023) == date(2023, 4, 15)
        assert week_code_to_date("W4", "April", 2023) == date(2023, 4, 22)
        assert week_code_to_date("W5", "April", 2023) == date(2023, 4, 29)

    def test_w5_clamped_to_short_month(self):
        # Feb 2023 has 28 days; W5 would compute day 29 -> must clamp to 28
        assert week_code_to_date("W5", "February", 2023) == date(2023, 2, 28)

    def test_admin_tunable_band_size_and_offset(self):
        # AppSetting override: band=5 days, offset=2 -> W1=day2, W2=day7
        assert week_code_to_date("W1", "April", 2023, band_size_days=5, start_offset_days=2) == date(2023, 4, 2)
        assert week_code_to_date("W2", "April", 2023, band_size_days=5, start_offset_days=2) == date(2023, 4, 7)


class TestResolveFinancialYearCalendarYear:
    def test_april_through_december_is_first_fy_year(self):
        for month in ["April", "May", "October", "December"]:
            assert _resolve_financial_year_calendar_year(month, "26-27") == 2026

    def test_january_through_march_is_second_fy_year(self):
        for month in ["January", "February", "March"]:
            assert _resolve_financial_year_calendar_year(month, "26-27") == 2027


class TestIsFooterRow:
    def test_total_row_detected(self):
        assert is_footer_row([None, "Total", None, 42])

    def test_compliance_legend_detected(self):
        assert is_footer_row(["% Compliance", None])

    def test_normal_machine_row_not_footer(self):
        assert not is_footer_row(["1", "AC-01", "Compressor", "Bosch"])


# ---------------------------------------------------------------------------
# normalize_for_dedup - the near-duplicate rule the customer specifically
# locked down: punctuation/whitespace differences collapse, but ANY letter
# or digit difference means a different physical machine.
# ---------------------------------------------------------------------------

class TestNormalizeForDedup:
    def test_space_dash_dot_are_equivalent(self):
        a = normalize_for_dedup("AC-01", "Compressor")
        b = normalize_for_dedup("AC 01", "Compressor")
        c = normalize_for_dedup("AC.01", "Compressor")
        assert a == b == c

    def test_case_insensitive(self):
        assert normalize_for_dedup("ac-01", "compressor") == normalize_for_dedup("AC-01", "COMPRESSOR")

    def test_different_digit_is_a_different_machine(self):
        assert normalize_for_dedup("AC-01", "Compressor") != normalize_for_dedup("AC-02", "Compressor")

    def test_different_letter_is_a_different_machine(self):
        # letter O swapped for digit 0 - looks similar to a human, must NOT match
        assert normalize_for_dedup("AC-01", "x") != normalize_for_dedup("AC-O1", "x")

    def test_extra_special_characters_only_still_matches(self):
        assert normalize_for_dedup("AC_01!!", "Compressor") == normalize_for_dedup("AC 01", "Compressor")


# ---------------------------------------------------------------------------
# parse_pm_sheet - end-to-end against synthetic sheets modeled on the
# customer's real layout (Sr No | M/c No. | NAME | Manufacture | Spec |
# Location | Remarks | (P/A) x12, header at row 3, data from row 4).
# ---------------------------------------------------------------------------

def _make_sheet(rows, header_row=3):
    """rows: list of lists of cell values, written starting at row 4."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(header_row, 1, "Sr No")
    ws.cell(header_row, 2, "M/c No.")
    ws.cell(header_row, 3, "Name of Machine")
    ws.cell(header_row, 8, "April P")
    ws.cell(header_row, 9, "April A")
    for i, row in enumerate(rows):
        r = header_row + 1 + i
        for c, val in enumerate(row, start=1):
            ws.cell(r, c, val)
    return ws


class TestParsePmSheet:
    def test_exact_date_plan_parsed(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "8.04.2023"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        assert len(result.rows) == 1
        plan = result.rows[0].plans[0]
        assert plan.planned_date == date(2023, 4, 8)
        assert plan.plan_source_type == "EXACT_DATE"

    def test_week_code_plan_parsed(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "W2"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        plan = result.rows[0].plans[0]
        assert plan.planned_date == date(2023, 4, 8)
        assert plan.planned_week == "W2"

    def test_day_of_month_plan_parsed(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "19"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        plan = result.rows[0].plans[0]
        assert plan.planned_date == date(2023, 4, 19)

    def test_day_of_month_exceeding_days_in_month_is_an_error_not_a_guess(self):
        # April has 30 days - "31" must error, never silently clamp/roll over
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "31"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        assert result.rows[0].plans == []
        assert any("31" in e for e in result.errors)

    def test_unknown_plan_value_skipped_with_warning_not_guessed(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "TBD"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        assert result.rows[0].plans == []
        assert any("unrecognized" in w.lower() for w in result.warnings)

    def test_copy_forward_p_equals_a_flagged_low_confidence(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "8.04.2023", "8.04.2023"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        plan = result.rows[0].plans[0]
        assert plan.low_confidence_actual is True

    def test_genuinely_different_actual_not_low_confidence(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "8.04.2023", "10.04.2023"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        plan = result.rows[0].plans[0]
        assert plan.low_confidence_actual is False

    def test_missing_machine_number_falls_back_to_name(self):
        # "Auto Phospating Plant" style row: no code, name only
        ws = _make_sheet([[1, None, "Auto Phospating Plant", None, None, None, None, "W1"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        row = result.rows[0]
        assert row.machine_number == "Auto Phospating Plant"
        assert any("fallback" in w.lower() for w in result.warnings)
        assert row.needs_review is True
        assert "machine_number_missing_used_name_fallback" in row.review_reasons

    def test_prose_in_machine_number_column_flagged(self):
        ws = _make_sheet([[1, "Auto Phospating Plant Section One Line", None, None, None, None, None, "W1"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        row = result.rows[0]
        assert row.needs_review is True
        assert "machine_number_looks_like_text" in row.review_reasons

    def test_year_in_remarks_flagged(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, "2019", "W1"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        row = result.rows[0]
        assert row.needs_review is True
        assert "remarks_looks_like_year" in row.review_reasons

    def test_ordinary_remarks_not_flagged(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, "runs fine", "W1"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        row = result.rows[0]
        assert "remarks_looks_like_year" not in row.review_reasons

    def test_date_outside_financial_year_flagged(self):
        # FY 23-24 means April=2023; an exact date typed with a stray 2022
        # year should be flagged, not silently accepted.
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "8.04.2022"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        plan = result.rows[0].plans[0]
        assert plan.needs_review is True
        assert "date_outside_financial_year" in plan.review_reasons
        assert result.rows[0].needs_review is True

    def test_critical_machine_flagged_by_asterisk(self):
        ws = _make_sheet([[1, "PP-01", "Power Press*", None, None, None, None, "W1"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        assert result.rows[0].critical is True

    def test_footer_row_stops_parsing(self):
        ws = _make_sheet([
            [1, "AC-01", "Compressor", None, None, None, None, "W1"],
            ["Total", None, None, None, None, None, None, None],
            [2, "AC-02", "Compressor", None, None, None, None, "W1"],
        ])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        # row after "Total" must never be parsed
        assert len(result.rows) == 1
        assert result.rows[0].machine_number == "AC-01"

    def test_completely_blank_row_is_skipped_silently(self):
        ws = _make_sheet([
            [1, "AC-01", "Compressor", None, None, None, None, "W1"],
            [None, None, None, None, None, None, None, None],
            [2, "AC-02", "Compressor", None, None, None, None, "W1"],
        ])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        assert len(result.rows) == 2

    def test_dedup_key_computed_on_every_row(self):
        ws = _make_sheet([[1, "AC-01", "Compressor", None, None, None, None, "W1"]])
        result = parse_pm_sheet(ws, "Sheet1", "23-24")
        assert result.rows[0].dedup_key == normalize_for_dedup("AC-01", "Compressor")
