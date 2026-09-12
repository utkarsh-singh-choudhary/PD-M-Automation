"""
Excel -> normalized PM plan importer.

Handles the three plan-encoding formats found in real company workbooks:
  1. EXACT_DATE     e.g. "8.04.2023"  (dd.mm.yyyy, dot or slash separated)
  2. WEEK_CODE      e.g. "W1".."W5"   -> resolved to a calendar date band
  3. DAY_OF_MONTH   e.g. "19"         -> the 19th of the listed month

Locked business rules (confirmed with the customer):
  - W1 = 1-7, W2 = 8-14, W3 = 15-21, W4 = 22-28, W5 = 29-end of month.
    The "planned_date" stored for a week code is the FIRST DAY of that band
    (reminders/escalations are computed relative to this date), while the
    original band is preserved in `planned_week` for display.
  - A day-of-month number N means "the Nth of that month".
  - Never silently invent dates: any cell that doesn't match one of the three
    known formats is recorded as a WARNING and skipped, not guessed.
  - If Plan and Actual cells for a given month are identical (copy-forward
    pattern seen in PM Plan-26-27), we still import it as Planned, but flag
    `low_confidence_actual=True` rather than treating it as a verified
    completion. A real completion must come from a user-entered PMActual.
"""

import calendar
import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Optional

import openpyxl

MONTH_HEADER_ORDER = [
    "April", "May", "June", "July", "August", "September",
    "October", "November", "December", "January", "February", "March",
]
MONTH_ABBR_MAP = {
    "apr": "April", "may": "May", "jun": "June", "jul": "July",
    "aug": "August", "sep": "September", "oct": "October",
    "nov": "November", "dec": "December", "jan": "January",
    "feb": "February", "mar": "March",
}

WEEK_BAND_SIZE_DAYS = 7  # locked convention: W1=1-7, W2=8-14, W3=15-21, W4=22-28, W5=29-end

DATE_PATTERNS = [
    re.compile(r"^\s*(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})\s*$"),
]
WEEK_PATTERN = re.compile(r"^\s*W([1-5])\s*$", re.IGNORECASE)
DAY_NUM_PATTERN = re.compile(r"^\s*(\d{1,2})\s*$")

# Rows/sections that are never machine data, regardless of which PM sheet.
FOOTER_MARKERS = {"total", "% compliance", "review", "* - for critical machine"}

# A 4-digit standalone number in Remarks is almost always a year (e.g. someone
# typed "2019" meaning "installed in 2019" or "AMC due 2024") that got left in
# a free-text column, not a real remark - worth a human glance, not an error.
YEAR_IN_REMARKS_PATTERN = re.compile(r"(?<!\d)(19|20)\d{2}(?!\d)")

# A "machine number" is expected to be a short code (letters/digits/./-),
# not prose. If it's long and contains multiple words, it's almost certainly
# a machine *name* that ended up in the number column (or vice versa).
LOOKS_LIKE_PROSE_PATTERN = re.compile(r"^[A-Za-z ]{15,}$")


def normalize_for_dedup(machine_number: str, machine_name: str) -> str:
    """
    Collapse a (machine_number, machine_name) pair down to just its letters
    and digits, lowercased, for near-duplicate detection.

    Deliberately narrow: this strips ONLY whitespace and punctuation/special
    characters (spaces, dots, dashes, underscores, etc.) so "AC-01" / "AC 01"
    / "AC.01" collapse to the same key. It does NOT do fuzzy/edit-distance
    matching - if any letter or digit differs ("AC-01" vs "AC-02", or
    "AC-01" vs "AC-O1"), the keys differ and the two rows are treated as
    different physical machines, per locked business rule: only punctuation/
    whitespace differences count as "the same machine, entered differently".
    """
    raw = f"{machine_number or ''}{machine_name or ''}".lower()
    return re.sub(r"[^a-z0-9]", "", raw)


@dataclass
class ParsedPlanRow:
    sr_no: Optional[int]
    machine_number: str
    machine_name: str
    manufacturer: Optional[str]
    specification: Optional[str]
    location: Optional[str]
    remarks: Optional[str]
    critical: bool
    plans: list = field(default_factory=list)  # list of ParsedMonthPlan
    needs_review: bool = False
    review_reasons: list = field(default_factory=list)  # short machine-level reason codes
    dedup_key: str = ""  # see normalize_for_dedup()


@dataclass
class ParsedMonthPlan:
    month: str
    plan_source_type: str
    plan_source_raw_value: str
    planned_date: Optional[date]
    planned_week: Optional[str]
    actual_raw_value: Optional[str]
    low_confidence_actual: bool
    source_cell_plan: str
    source_cell_actual: str
    needs_review: bool = False
    review_reasons: list = field(default_factory=list)


@dataclass
class ImportResult:
    rows: list = field(default_factory=list)   # ParsedPlanRow
    warnings: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    sheet_name: str = ""
    financial_year: str = ""


def compute_file_hash(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def week_code_to_date(week_code: str, month_name: str, year: int,
                       band_size_days: int = None, start_offset_days: int = None) -> date:
    """W1='1'->day 1, W2->day 8, W3->day15, W4->day22, W5->day29 (clamped to month length).

    band_size_days / start_offset_days let this be retuned from the admin panel
    (AppSetting keys week_band_size_days / week_start_offset_days) without a
    code change; both default to the locked convention (7, 1) when omitted.
    """
    week_num = int(WEEK_PATTERN.match(week_code).group(1))
    month_num = MONTH_HEADER_ORDER.index(month_name) # not calendar month number directly
    # translate month name to actual calendar month number (1-12)
    calendar_month = _month_name_to_number(month_name)
    last_day = calendar.monthrange(year, calendar_month)[1]
    band_size = band_size_days if band_size_days is not None else WEEK_BAND_SIZE_DAYS
    start_offset = start_offset_days if start_offset_days is not None else 1
    day = min(start_offset + (week_num - 1) * band_size, last_day)
    return date(year, calendar_month, day)


def _month_name_to_number(month_name: str) -> int:
    return list(calendar.month_name).index(month_name)


FY_PATTERN = re.compile(r"^\d{2}-\d{2}$")


def validate_financial_year(fy_label: str) -> None:
    """Raises ValueError with a clear message if fy_label isn't the
    'YY-YY' format every other part of the importer assumes (Medium: "Excel
    import needs stronger validation around commit" - previously a
    malformed financial_year caused an unhandled 500 deep inside
    _resolve_financial_year_calendar_year's `fy_label.split('-')`)."""
    if not FY_PATTERN.match(fy_label or ""):
        raise ValueError(
            f"Invalid financial_year '{fy_label}' - expected format 'YY-YY' (e.g. '25-26')"
        )
    start_yy, end_yy = fy_label.split("-")
    if int(end_yy) != (int(start_yy) + 1) % 100:
        raise ValueError(
            f"Invalid financial_year '{fy_label}' - the two years must be consecutive (e.g. '25-26', not '25-28')"
        )


def _resolve_financial_year_calendar_year(month_name: str, fy_label: str) -> int:
    """
    fy_label like '26-27' means Apr 2026 - Mar 2027.
    April-December -> first year of FY; Jan-March -> second year of FY.
    """
    validate_financial_year(fy_label)
    start_yy, end_yy = fy_label.split("-")
    start_year = 2000 + int(start_yy)
    end_year = 2000 + int(end_yy)
    if month_name in ("January", "February", "March"):
        return end_year
    return start_year


def _parse_exact_date(raw: str) -> Optional[date]:
    for pat in DATE_PATTERNS:
        m = pat.match(raw)
        if m:
            d, mo, y = m.groups()
            y = int(y)
            if y < 100:
                y += 2000
            # Sanity bound (Medium: "Excel import needs stronger validation
            # around commit" - a fat-fingered year like "2 02 5" parsing to
            # 2002 or 0205, or a stray "9999", shouldn't silently become a
            # PM plan date decades away). PM plans are always within a
            # single financial year, so anything this far outside "now" is
            # certainly a data entry error, not a real maintenance date.
            if not (2015 <= y <= 2100):
                return None
            try:
                return date(y, int(mo), int(d))
            except ValueError:
                return None
    return None


def classify_cell_value(value) -> tuple:
    """Returns (type, raw_str) where type in EXACT_DATE / WEEK_CODE / DAY_OF_MONTH / UNKNOWN / EMPTY."""
    if value is None or value == "":
        return "EMPTY", None
    if isinstance(value, (datetime, date)):
        return "EXACT_DATE", value.strftime("%d.%m.%Y")
    raw = str(value).strip()
    if WEEK_PATTERN.match(raw):
        return "WEEK_CODE", raw.upper()
    if DATE_PATTERNS[0].match(raw):
        return "EXACT_DATE", raw
    if DAY_NUM_PATTERN.match(raw):
        return "DAY_OF_MONTH", raw
    return "UNKNOWN", raw


def is_footer_row(row_values: list) -> bool:
    joined = " ".join(str(v).strip().lower() for v in row_values if v is not None)
    return any(marker in joined for marker in FOOTER_MARKERS)


def parse_pm_sheet(
    ws,
    sheet_name: str,
    financial_year: str,
    header_row: int = 3,
    data_start_row: int = 4,
    col_map: Optional[dict] = None,
    week_band_size_days: int = None,
    week_start_offset_days: int = None,
) -> ImportResult:
    """
    Parses a PM/PD plan sheet that follows the "Sr No | M/c No. | NAME OF MACHINE |
    Manufacture | SPECIFICATION | Location | REMARKS | (P/A)x12" layout, with either
    a leading "Sr No" column or not (col_map lets the caller override positions).
    """
    result = ImportResult(sheet_name=sheet_name, financial_year=financial_year)

    col_map = col_map or {
        "sr_no": 1, "machine_number": 2, "machine_name": 3,
        "manufacturer": 4, "specification": 5, "location": 6, "remarks": 7,
        "first_month_col": 8,  # first "P" column; A is first_month_col+1, next month +2, etc.
    }

    max_row = ws.max_row
    max_col = ws.max_column

    for r in range(data_start_row, max_row + 1):
        row_values = [ws.cell(r, c).value for c in range(1, max_col + 1)]
        if all(v is None for v in row_values):
            continue
        if is_footer_row(row_values):
            break  # footer/legend section reached; stop parsing data rows

        machine_number = row_values[col_map["machine_number"] - 1]
        machine_name = row_values[col_map["machine_name"] - 1]

        if not machine_number and not machine_name:
            result.warnings.append(f"Row {r}: skipped, no machine number or name.")
            continue
        if not machine_number:
            # e.g. "Auto Phospating Plant" style rows where name is in the number column
            machine_number = str(machine_name)
            result.warnings.append(
                f"Row {r}: machine number missing, used machine name as fallback identifier."
            )

        remarks_clean = _clean(row_values[col_map["remarks"] - 1])
        row_reasons = []

        # machine number missing entirely (fallback already applied above)
        if not row_values[col_map["machine_number"] - 1]:
            row_reasons.append("machine_number_missing_used_name_fallback")
        # machine-number column holds prose instead of a short code - likely
        # swapped with the name column, or the sheet has no real code for
        # this row at all.
        elif LOOKS_LIKE_PROSE_PATTERN.match(str(machine_number).strip()):
            row_reasons.append("machine_number_looks_like_text")
        # a bare 4-digit year sitting in "Remarks" (installation/AMC year
        # typed into the wrong field) rather than an actual remark.
        if remarks_clean and YEAR_IN_REMARKS_PATTERN.search(remarks_clean) and len(remarks_clean) <= 6:
            row_reasons.append("remarks_looks_like_year")

        row = ParsedPlanRow(
            sr_no=row_values[col_map["sr_no"] - 1] if col_map.get("sr_no") else None,
            machine_number=str(machine_number).strip(),
            machine_name=str(machine_name).strip() if machine_name else str(machine_number).strip(),
            manufacturer=_clean(row_values[col_map["manufacturer"] - 1]),
            specification=_clean(row_values[col_map["specification"] - 1]),
            location=_clean(row_values[col_map["location"] - 1]),
            remarks=remarks_clean,
            critical="*" in str(machine_name or ""),
            needs_review=bool(row_reasons),
            review_reasons=row_reasons,
            dedup_key=normalize_for_dedup(str(machine_number).strip(), str(machine_name or "").strip()),
        )

        col = col_map["first_month_col"]
        for month_name in MONTH_HEADER_ORDER:
            if col > max_col:
                break
            plan_val = ws.cell(r, col).value
            actual_val = ws.cell(r, col + 1).value if col + 1 <= max_col else None

            plan_type, plan_raw = classify_cell_value(plan_val)
            if plan_type in ("EMPTY",):
                col += 2
                continue
            if plan_type == "UNKNOWN":
                result.warnings.append(
                    f"Row {r} ({row.machine_number}) {month_name}: unrecognized plan value "
                    f"'{plan_val}' - skipped, not imported. Needs manual review."
                )
                col += 2
                continue

            calendar_year = _resolve_financial_year_calendar_year(month_name, financial_year)
            planned_date, planned_week = None, None

            try:
                if plan_type == "EXACT_DATE":
                    planned_date = _parse_exact_date(plan_raw) if isinstance(plan_val, str) else plan_val.date() if isinstance(plan_val, datetime) else plan_val
                elif plan_type == "WEEK_CODE":
                    planned_date = week_code_to_date(plan_raw, month_name, calendar_year,
                                                      band_size_days=week_band_size_days,
                                                      start_offset_days=week_start_offset_days)
                    planned_week = plan_raw
                elif plan_type == "DAY_OF_MONTH":
                    day_num = int(plan_raw)
                    month_num = _month_name_to_number(month_name)
                    last_day = calendar.monthrange(calendar_year, month_num)[1]
                    if day_num > last_day:
                        raise ValueError(f"day {day_num} exceeds days in {month_name}")
                    planned_date = date(calendar_year, month_num, day_num)
            except (ValueError, AttributeError) as e:
                result.errors.append(
                    f"Row {r} ({row.machine_number}) {month_name}: failed to convert "
                    f"'{plan_val}' -> {e}"
                )
                col += 2
                continue

            actual_type, actual_raw = classify_cell_value(actual_val)
            low_confidence = actual_type != "EMPTY" and actual_raw == plan_raw

            plan_reasons = []
            if planned_date is not None:
                expected_year = _resolve_financial_year_calendar_year(month_name, financial_year)
                if planned_date.year != expected_year:
                    plan_reasons.append("date_outside_financial_year")

            row.plans.append(ParsedMonthPlan(
                month=month_name,
                plan_source_type=plan_type,
                plan_source_raw_value=str(plan_val),
                planned_date=planned_date,
                planned_week=planned_week,
                actual_raw_value=str(actual_val) if actual_val is not None else None,
                low_confidence_actual=low_confidence,
                source_cell_plan=ws.cell(r, col).coordinate,
                source_cell_actual=ws.cell(r, col + 1).coordinate if col + 1 <= max_col else "",
                needs_review=bool(plan_reasons),
                review_reasons=plan_reasons,
            ))
            if plan_reasons:
                row.needs_review = True
                row.review_reasons = row.review_reasons + [f"{month_name}:{r}" for r in plan_reasons]
            col += 2

        result.rows.append(row)

    return result


def _clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def load_workbook_from_path(path: str):
    return openpyxl.load_workbook(path, data_only=True)


def load_workbook_from_bytes(data: bytes):
    """Preferred over load_workbook_from_path now that uploads live in
    object storage, not on a local filesystem path - avoids needing a
    scratch file on disk just to parse an already-in-memory workbook."""
    import io
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
